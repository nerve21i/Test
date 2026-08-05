#!/usr/bin/env python3
"""PDF 파일을 OCR로 읽어 엑셀(.xlsx) 파일로 변환하는 CLI 도구."""

import argparse
import sys
from dataclasses import dataclass

try:
    import fitz  # PyMuPDF
except ImportError:
    sys.exit("PyMuPDF가 설치되어 있지 않습니다. 'pip install -r requirements.txt' 를 실행하세요.")

try:
    import pytesseract
    from pytesseract import Output
except ImportError:
    sys.exit("pytesseract가 설치되어 있지 않습니다. 'pip install -r requirements.txt' 를 실행하세요.")

try:
    from PIL import Image
except ImportError:
    sys.exit("Pillow가 설치되어 있지 않습니다. 'pip install -r requirements.txt' 를 실행하세요.")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl이 설치되어 있지 않습니다. 'pip install -r requirements.txt' 를 실행하세요.")


@dataclass
class LineResult:
    page: int
    block: int
    paragraph: int
    line: int
    text: str
    confidence: float


def check_tesseract():
    try:
        pytesseract.get_tesseract_version()
    except Exception:
        sys.exit(
            "Tesseract OCR 엔진을 찾을 수 없습니다.\n"
            "설치 방법 (Ubuntu/Debian): sudo apt-get install tesseract-ocr tesseract-ocr-kor\n"
            "설치 후 다시 실행해주세요."
        )


def page_has_text(page) -> bool:
    return bool(page.get_text("text").strip())


def extract_embedded_text(page) -> list[LineResult]:
    results = []
    raw = page.get_text("text")
    for i, line in enumerate(l for l in raw.splitlines() if l.strip()):
        results.append(LineResult(page=page.number + 1, block=0, paragraph=0, line=i, text=line.strip(), confidence=100.0))
    return results


def render_page_to_image(page, dpi: int) -> Image.Image:
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=matrix, alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def ocr_page(page, dpi: int, lang: str) -> list[LineResult]:
    image = render_page_to_image(page, dpi)
    data = pytesseract.image_to_data(image, lang=lang, output_type=Output.DICT)

    lines: dict[tuple[int, int, int], dict] = {}
    n = len(data["text"])
    for i in range(n):
        word = data["text"][i].strip()
        if not word:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        conf = data["conf"][i]
        try:
            conf = float(conf)
        except ValueError:
            conf = -1.0
        entry = lines.setdefault(key, {"words": [], "top": data["top"][i], "confs": []})
        entry["words"].append(word)
        if conf >= 0:
            entry["confs"].append(conf)

    ordered_keys = sorted(lines.keys(), key=lambda k: (k[0], k[1], lines[k]["top"]))
    results = []
    for block, par, line_no in ordered_keys:
        entry = lines[(block, par, line_no)]
        text = " ".join(entry["words"])
        confs = entry["confs"]
        avg_conf = round(sum(confs) / len(confs), 1) if confs else 0.0
        results.append(LineResult(page=page.number + 1, block=block, paragraph=par, line=line_no, text=text, confidence=avg_conf))
    return results


def process_pdf(pdf_path: str, lang: str, dpi: int, mode: str, progress=print) -> list[LineResult]:
    doc = fitz.open(pdf_path)
    all_results: list[LineResult] = []

    for page in doc:
        page_num = page.number + 1
        use_ocr = mode == "ocr" or (mode == "auto" and not page_has_text(page))

        if mode == "text" or (mode == "auto" and not use_ocr):
            progress(f"[{page_num}/{doc.page_count}] 페이지: 내장 텍스트 추출")
            page_results = extract_embedded_text(page)
        else:
            progress(f"[{page_num}/{doc.page_count}] 페이지: OCR 처리 중...")
            page_results = ocr_page(page, dpi, lang)

        all_results.extend(page_results)

    doc.close()
    return all_results


def write_excel(results: list[LineResult], output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "OCR 결과"

    headers = ["페이지", "블록", "문단", "줄", "텍스트", "신뢰도(%)"]
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)

    for r in results:
        ws.append([r.page, r.block, r.paragraph, r.line, r.text, r.confidence])

    widths = [8, 8, 8, 8, 80, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="PDF 파일을 OCR로 읽어 엑셀 파일로 변환합니다.")
    parser.add_argument("pdf", help="입력 PDF 파일 경로")
    parser.add_argument("-o", "--output", help="출력 엑셀 파일 경로 (기본: 입력 파일명 + .xlsx)")
    parser.add_argument("--lang", default="kor+eng", help="OCR 언어 (기본: kor+eng). 예: eng, kor, kor+eng")
    parser.add_argument("--dpi", type=int, default=300, help="OCR용 이미지 해상도 (기본: 300)")
    parser.add_argument(
        "--mode",
        choices=["auto", "ocr", "text"],
        default="auto",
        help="auto: 텍스트가 있으면 추출, 없으면 OCR(기본) / ocr: 모든 페이지 강제 OCR / text: 내장 텍스트만 추출",
    )
    args = parser.parse_args()

    if args.mode != "text":
        check_tesseract()

    output_path = args.output or (args.pdf.rsplit(".", 1)[0] + ".xlsx")

    print(f"PDF 처리 시작: {args.pdf}")
    results = process_pdf(args.pdf, args.lang, args.dpi, args.mode)

    if not results:
        print("경고: 추출된 텍스트가 없습니다.")

    write_excel(results, output_path)
    print(f"완료: {len(results)}개 줄을 '{output_path}' 파일로 저장했습니다.")


if __name__ == "__main__":
    main()
